# =========================
# STANDARD / THIRD-PARTY IMPORTS
# =========================
import streamlit as st
import requests
import json
import numpy as np
from io import BytesIO
from typing import List, Dict

# =========================
# LOCAL PROJECT IMPORTS
# =========================
# These callbacks manage UI state changes (button clicks, toggles, chat loading)
from callbacks import (
    handle_new_chat_click,
    toggle_upload,
    toggle_mic,
    load_chat
)

# Injects custom CSS styles for chat bubbles, layout, footer, etc.
from styles import inject_css


# =========================
# APPLICATION CONFIGURATION
# =========================
class AppConfig:
    """
    Centralized configuration class.
    Keeping constants here avoids magic values spread across the code.
    """
    PAGE_TITLE = "Code Genie"
    PAGE_ICON = "Logo-removebg-preview.png"
    LOGO_IMAGE = "Logo1-removebg-preview.png"
    GITHUB_LINK = "https://github.com/manojkarri04/CodeGenie"

    # Ollama local API endpoint
    OLLAMA_URL = "http://localhost:11434/api/chat"
    MODEL_NAME = "CodeGenie"

    # Used to downscale very large images before OCR (performance optimization)
    MAX_IMG_WIDTH = 1000

    # Voice recognition constraints
    VOICE_TIMEOUT = 5
    VOICE_PHRASE_LIMIT = 10


# =========================
# PAGE & SESSION INITIALIZATION
# =========================
def configure_page():
    """
    Sets Streamlit page-level settings and displays the app logo.
    """
    st.set_page_config(
        page_title=AppConfig.PAGE_TITLE,
        page_icon=AppConfig.PAGE_ICON,
        layout="centered",
        initial_sidebar_state="expanded"
    )

    # st.logo is available only in newer Streamlit versions,
    # so hasattr is used to avoid crashes in older versions
    if hasattr(st, "logo"):
        st.logo(
            AppConfig.LOGO_IMAGE,
            link=AppConfig.GITHUB_LINK,
            icon_image=AppConfig.PAGE_ICON
        )


def init_session_state():
    """
    Initializes all required session_state variables.
    This prevents KeyError and ensures predictable app behavior.
    """
    defaults = {
        "messages": [],            # Stores chat messages (user + assistant)
        "chat_started": False,     # Used to control UI layout
        "upload_clicked": False,   # Controls file uploader visibility
        "uploaded_file": None,     # Stores uploaded file object
        "mic_active": False,       # Mic button toggle state
        "input_key": 0,            # Used to reset text_area input
        "recent_chats": [],        # Sidebar chat history
        "listening": False,        # Voice listening state
        "audio": None              # Temporarily stores recorded audio
    }

    for key, val in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = val


# =========================
# OCR SERVICE
# =========================
@st.cache_resource
def load_ocr_reader():
    """
    Loads EasyOCR reader once and caches it.
    Heavy initialization → cache_resource improves performance.
    """
    import easyocr
    return easyocr.Reader(['en'], gpu=False)


@st.cache_data(show_spinner=True)
def process_file_ocr(file_bytes: BytesIO, file_type: str) -> str:
    """
    Extracts text from uploaded files.
    Supports:
    - Images (OCR)
    - PDF
    - DOCX
    - PPTX
    - TXT

    Returns extracted text as a single string.
    """
    try:
        # ---------- IMAGE OCR ----------
        if file_type.startswith("image/"):
            from PIL import Image

            reader = load_ocr_reader()

            # Convert image to grayscale → improves OCR accuracy
            image = Image.open(file_bytes).convert("L")

            # Resize large images to reduce OCR computation time
            if image.width > AppConfig.MAX_IMG_WIDTH:
                ratio = AppConfig.MAX_IMG_WIDTH / image.width
                image = image.resize(
                    (AppConfig.MAX_IMG_WIDTH, int(image.height * ratio))
                )

            # detail=0 returns only text (no bounding boxes)
            text_list = reader.readtext(
                np.array(image),
                detail=0,
                paragraph=True
            )
            return " ".join(text_list)

        # ---------- PDF ----------
        elif file_type == "application/pdf":
            from pypdf import PdfReader

            reader = PdfReader(file_bytes)
            return "\n".join(
                page.extract_text()
                for page in reader.pages
                if page.extract_text()
            )

        # ---------- WORD DOCUMENT ----------
        elif "wordprocessingml" in file_type:
            from docx import Document

            doc = Document(file_bytes)
            return "\n".join(
                p.text for p in doc.paragraphs if p.text.strip()
            )

        # ---------- POWERPOINT ----------
        elif "presentation" in file_type:
            from pptx import Presentation

            prs = Presentation(file_bytes)
            text = []

            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text.append(shape.text)

            return "\n".join(text)

        # ---------- PLAIN TEXT ----------
        elif file_type == "text/plain":
            return file_bytes.read().decode("utf-8")

        else:
            st.warning(f"Unsupported file type: {file_type}")
            return ""

    except Exception as e:
        st.error(f"Error processing file: {e}")
        return ""


# =========================
# VOICE INPUT SERVICE
# =========================
def process_voice_input() -> str:
    """
    Handles microphone input and speech-to-text conversion.
    Uses Google Speech Recognition via SpeechRecognition library.
    """
    import speech_recognition as sr

    recognizer = sr.Recognizer()

    with sr.Microphone() as source:

        # First phase: listening
        if st.session_state.listening:
            st.info("🎙️ Listening... Speak now!")
            try:
                audio = recognizer.listen(
                    source,
                    timeout=AppConfig.VOICE_TIMEOUT,
                    phrase_time_limit=AppConfig.VOICE_PHRASE_LIMIT
                )
                st.session_state.audio = audio
            except sr.WaitTimeoutError:
                st.warning("❌ No speech detected.")
                return ""

        # Second phase: processing captured audio
        else:
            if st.session_state.get("audio"):
                st.info("⏳ Processing your voice...")
                audio = st.session_state.audio
                st.session_state.audio = None

                try:
                    text = recognizer.recognize_google(audio)
                    st.success(f"✅ You said: {text}")
                    return text
                except sr.UnknownValueError:
                    st.warning("❌ Couldn't understand speech.")
                except sr.RequestError as e:
                    st.error(f"⚠️ Speech service error: {e}")

    return ""


# =========================
# OLLAMA API STREAMING
# =========================
def stream_ollama_response(messages: List[Dict[str, str]]) -> str:
    """
    Sends chat history to Ollama and streams the response token-by-token.
    UI updates in real-time using a placeholder.
    """
    payload = {
        "model": AppConfig.MODEL_NAME,
        "messages": messages,
        "stream": True
    }

    placeholder = st.empty()
    full_text = ""

    try:
        with requests.post(
            AppConfig.OLLAMA_URL,
            json=payload,
            stream=True
        ) as response:

            if response.status_code == 200:
                for line in response.iter_lines():
                    if line:
                        body = json.loads(line)

                        if "message" in body and "content" in body["message"]:
                            full_text += body["message"]["content"]
                            _render_streaming_bubble(placeholder, full_text)

                        if body.get("done"):
                            break
            else:
                full_text = "Error: Could not connect to CodeGenie."

    except Exception as e:
        full_text = f"Error: {e}"

    # Final render (without cursor)
    placeholder.html(f"""
        <div class="chat-message-row ai-message-row">
            <div class="chat-message-bubble ai-message-bubble">{full_text}</div>
        </div>
    """)

    return full_text


def _render_streaming_bubble(placeholder, text: str):
    """
    Displays intermediate streaming text with a cursor effect.
    """
    placeholder.markdown(f"""
        <div class="chat-message-row ai-message-row">
            <div class="chat-message-bubble ai-message-bubble">{text}▌</div>
        </div>
    """, unsafe_allow_html=True)


# =========================
# UI RENDERING
# =========================
def render_sidebar():
    """
    Sidebar with New Chat button and recent chats.
    """
    with st.sidebar:
        st.button("➕ New Chat", on_click=handle_new_chat_click, use_container_width=True)
        st.divider()
        st.markdown("<h3>Recent chats</h3>", unsafe_allow_html=True)

        for i, chat in enumerate(st.session_state.recent_chats):
            st.button(
                chat["title"],
                on_click=load_chat,
                args=(i,),
                key=f"chat_{i}",
                use_container_width=True
            )

        # Profile footer
        st.html("""
            <div class="profile-container">
                <div class="avatar"><b></b></div>
                <div class="details"><b>Code Genie</b></div>
            </div>
        """)


def render_chat_history():
    """
    Displays last 40 chat messages.
    """
    container_cls = (
        "chat-history-container visible"
        if st.session_state.chat_started
        else "chat-history-container"
    )

    st.markdown(f'<div class="{container_cls}">', unsafe_allow_html=True)

    for msg in st.session_state.messages[-40:]:
        role = msg["role"]
        row_cls = "user-message-row" if role == "user" else "ai-message-row"
        bubble_cls = "user-message-bubble" if role == "user" else "ai-message-bubble"

        st.markdown(f"""
            <div class="chat-message-row {row_cls}">
                <div class="chat-message-bubble {bubble_cls}">
                    {msg["content"]}
                </div>
            </div>
        """, unsafe_allow_html=True)

    st.markdown('<div style="height:100px;"></div>', unsafe_allow_html=True)
    st.markdown("</div>", unsafe_allow_html=True)


def render_file_uploader():
    """
    Shows file uploader only when upload button is clicked.
    """
    if st.session_state.upload_clicked and st.session_state.uploaded_file is None:
        placeholder = st.empty()

        with placeholder.container():
            st.markdown('<div class="fixed-footer" style="bottom:80px;">', unsafe_allow_html=True)
            uploaded = st.file_uploader("Upload your file", type=["pdf", "png", "jpg", "txt"])
            st.markdown('</div>', unsafe_allow_html=True)

        if uploaded:
            st.session_state.uploaded_file = uploaded
            placeholder.empty()

            text = process_file_ocr(uploaded, uploaded.type)
            if text:
                _handle_user_submission(f"📎 Extracted text:\n\n{text[:500]}")

            st.session_state.upload_clicked = False
            st.rerun()


def render_input_area():
    """
    Bottom input bar with text area, upload button, and mic.
    """
    footer_cls = "fixed-footer" if st.session_state.chat_started else "search-bar-container"
    st.markdown(f'<div class="{footer_cls}">', unsafe_allow_html=True)

    col1, col2, col3 = st.columns([1, 10, 1])

    with col1:
        st.button("➕", on_click=toggle_upload)

    with col2:
        user_input = st.text_area(
            "Message CodeGenie...",
            height=52,
            key=f"query_{st.session_state.input_key}",
            label_visibility="collapsed"
        )

    with col3:
        st.button("🎙️" if st.session_state.mic_active else "🎤", on_click=toggle_mic)

    # Voice handling
    if st.session_state.mic_active:
        st.session_state.listening = not st.session_state.listening
        voice_text = process_voice_input()
        if voice_text:
            user_input = voice_text
            st.session_state.mic_active = False

    st.markdown("</div>", unsafe_allow_html=True)

    if user_input and user_input.strip():
        _handle_user_submission(user_input)


def _handle_user_submission(content: str):
    """
    Handles user message submission and assistant response.
    """
    st.session_state.chat_started = True

    if "title_placeholder" in st.session_state:
        st.session_state.title_placeholder.empty()

    st.session_state.messages.append({"role": "user", "content": content})

    reply = stream_ollama_response(st.session_state.messages)
    st.session_state.messages.append({"role": "assistant", "content": reply})

    st.session_state.input_key += 1
    st.rerun()


def _handle_user_interface():
    """
    Controls the main title visibility.
    """
    title_placeholder = st.empty()
    st.session_state.title_placeholder = title_placeholder

    if not st.session_state.chat_started:
        title_placeholder.markdown('<h1 class="title">What can I help with?</h1>', unsafe_allow_html=True)
    else:
        title_placeholder.empty()


# =========================
# MAIN ENTRY POINT
# =========================
def main():
    configure_page()
    init_session_state()
    inject_css()
    _handle_user_interface()
    render_chat_history()
    render_file_uploader()
    render_sidebar()
    render_input_area()


if __name__ == "__main__":
    main()
